"""Unified DOCX, PPTX, XLSX, and XLS evidence parsing."""

from __future__ import annotations

import hashlib
import mimetypes
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from langchain_core.documents import Document

from app.services.evidence.models import (
    EvidenceDocument,
    ImageBlock,
    ParsedDocument,
    ParsedPage,
    TableBlock,
    TextBlock,
)

_DOCX = ".docx"
_PPTX = ".pptx"
_XLSX = ".xlsx"
_XLS = ".xls"
_CSV = ".csv"
OFFICE_EXTENSIONS = frozenset({_DOCX, _PPTX, _XLSX, _XLS, _CSV})

# The formats read straight out of the OOXML archive -- workbooks and delimited
# tables are routed to _load_workbook before either reader below is reached.
# A table rather than "docx, otherwise pptx", so a format added above raises here
# instead of being silently read as a slide deck.
_ARCHIVE_PREFIXES = {
    _DOCX: {"pages": "word/", "media": "word/media/"},
    _PPTX: {"pages": "ppt/slides/", "media": "ppt/media/"},
}


def load_office_document(path: Path, document: EvidenceDocument) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix not in OFFICE_EXTENSIONS:
        raise ValueError(f"unsupported Office format: {suffix}")
    if suffix in {_XLSX, _XLS, _CSV}:
        return _load_workbook(path, document)

    markdown = _docling_markdown(path)
    fallback_chain = ["docling"]
    parser = "docling"
    archive_pages = _archive_pages(path, suffix)
    if not markdown:
        markdown = "\n\n".join(archive_pages)
        fallback_chain.append("office_xml")
        parser = "office_xml"
    elif suffix == _PPTX and archive_pages:
        fallback_chain.append("office_xml_page_map")
    images = _archive_images(path, document)
    page_texts = archive_pages if suffix == _PPTX and archive_pages else [markdown]
    pages = tuple(ParsedPage(page=index, text=text) for index, text in enumerate(page_texts, start=1))
    blocks = tuple(
        TextBlock(block_id=_id(document, "text", index), page=index, text=text)
        for index, text in enumerate(page_texts, start=1)
        if text
    )
    docx_tables: tuple[TableBlock, ...] = ()
    if suffix == _DOCX and markdown:
        docx_tables = _extract_markdown_tables(markdown, document)
    return ParsedDocument(
        document=document,
        pages=pages,
        text_blocks=blocks,
        images=images,
        tables=docx_tables,
        parser=parser,
        fallback_chain=tuple(fallback_chain),
    )


def parsed_to_documents(parsed: ParsedDocument) -> list[Document]:
    base = {
        "source": parsed.document.source,
        "filename": parsed.document.filename,
        "document_id": parsed.document.document_id,
        "version": parsed.document.version,
        "tenant_id": parsed.document.tenant_id,
        "owner_user_id": parsed.document.owner_user_id,
        "visibility": parsed.document.visibility,
        "acl_tags": ",".join(parsed.document.acl_tags),
        "parser": parsed.parser,
    }
    # When text_blocks exist (e.g. DOCX), the markdown text already contains the tables inline.
    # Appending both text_blocks and parsed.tables causes duplicate documents and redundant chunking.
    # parsed.tables remains on ParsedDocument for specialized multimodal _index_tables.
    if parsed.text_blocks:
        documents = [
            Document(
                page_content=block.text,
                metadata={
                    **base,
                    "page": block.page,
                    "sheet": block.sheet or "",
                    "block_id": block.block_id,
                    "modality": "text",
                },
            )
            for block in parsed.text_blocks
            if block.text.strip()
        ]
    else:
        documents = [
            Document(
                page_content=table.markdown,
                metadata={
                    **base,
                    "page": table.page,
                    "sheet": table.sheet or "",
                    "table_id": table.table_id,
                    "modality": "table",
                },
            )
            for table in parsed.tables
            if table.markdown.strip()
        ]
    documents.extend(
        Document(
            page_content=(image.ocr_text or image.description or f"Image artifact {image.image_id}"),
            metadata={**base, "page": image.page, "image_id": image.image_id, "modality": "image"},
        )
        for image in parsed.images
    )
    return documents


def _load_workbook(path: Path, document: EvidenceDocument) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == _XLSX:
        sheets = _xlsx_rows(path)
        parser = "openpyxl"
    elif suffix == _XLS:
        sheets = _xls_rows(path)
        parser = "pandas"
    else:
        sheets = _csv_rows(path)
        parser = "csv"

    pages: list[ParsedPage] = []
    tables: list[TableBlock] = []
    for page_number, (sheet, rows) in enumerate(sheets, start=1):
        sheet_label = sheet if (sheet and sheet.strip()) else None
        markdown = _rows_to_markdown(rows, sheet_name=sheet_label)
        pages.append(ParsedPage(page=page_number, sheet=sheet, text=markdown))
        if markdown:
            tables.append(
                TableBlock(
                    table_id=_id(document, "table", page_number),
                    page=page_number,
                    sheet=sheet,
                    markdown=markdown,
                )
            )
    return ParsedDocument(
        document=document,
        pages=tuple(pages),
        tables=tuple(tables),
        parser=parser,
        fallback_chain=(parser,),
    )


MAX_TABLE_ROWS: int = 5000

# Full-mode openpyxl is what exposes merged ranges, and it also materialises one
# MergedCell per merged cell while loading -- so a few bytes of XML declaring
# A1:XFD1048576 exhaust memory inside openpyxl, before any guard of ours runs.
# The declared area is read from the archive first; past these limits the
# workbook is streamed (read-only) and merged cells are simply not filled.
_MAX_MERGED_CELLS = 200_000
_FULL_LOAD_MAX_BYTES = 20 * 1024 * 1024
_MAX_SHEET_XML_SCAN_BYTES = 256 * 1024 * 1024
_MERGE_REF_RE = re.compile(rb'<(?:\w+:)?mergeCell\b[^>]*?\bref="([A-Z]{1,3})(\d{1,7}):([A-Z]{1,3})(\d{1,7})"')


def _column_number(letters: bytes) -> int:
    number = 0
    for code in letters:
        number = number * 26 + (code - 64)
    return number


def _scan_sheet_merge_refs(handle: Any, scanned: int, total: int) -> tuple[int, int, bool]:
    """Scan one worksheet XML stream for merged ranges. Returns (scanned, total, is_over_limit)."""
    tail = b""
    while chunk := handle.read(1 << 20):
        scanned += len(chunk)
        if scanned > _MAX_SHEET_XML_SCAN_BYTES:
            return scanned, _MAX_MERGED_CELLS + 1, True
        buffer = tail + chunk
        last_end = 0
        for match in _MERGE_REF_RE.finditer(buffer):
            c1, r1, c2, r2 = match.groups()
            rows = abs(int(r2) - int(r1)) + 1
            cols = abs(_column_number(c2) - _column_number(c1)) + 1
            total += rows * cols
            if total > _MAX_MERGED_CELLS:
                return scanned, total, True
            last_end = match.end()
        tail = buffer[max(last_end, len(buffer) - 256) :]
    return scanned, total, False


def _declared_merged_area(path: Path) -> int:
    """Cells covered by merged ranges across every sheet, read from the raw XML.

    Streams each sheet so a decompression bomb cannot be materialised here
    either; anything unreadable or oversized reports "over the limit", which
    only means the workbook is read in streaming mode.
    """
    over_limit = _MAX_MERGED_CELLS + 1
    total = 0
    scanned = 0
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                    continue
                with archive.open(name) as handle:
                    scanned, total, is_over = _scan_sheet_merge_refs(handle, scanned, total)
                    if is_over:
                        return total
    except (zipfile.BadZipFile, OSError, KeyError):
        return over_limit
    return total


def _csv_rows(path: Path) -> list[tuple[str, list[list[object]]]]:
    import csv

    raw = path.read_bytes()
    text = ""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "latin1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        text = raw.decode("latin1", errors="ignore")

    sample = text[:4096]
    delimiter = ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except Exception:
        if "\t" in sample and sample.count("\t") > sample.count(","):
            delimiter = "\t"

    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    rows: list[list[object]] = [list(row) for row in reader if any(str(cell).strip() for cell in row)]
    total_rows = len(rows)
    if total_rows > MAX_TABLE_ROWS:
        rows = rows[:MAX_TABLE_ROWS]
        rows.append([f"... [Truncated: {total_rows} total rows, first {MAX_TABLE_ROWS} rows shown] ..."])

    sheet_name = path.stem or "Sheet1"
    return [(sheet_name, rows)]


def _extract_sheet_rows_with_merged_cells(sheet: Any, max_rows: int = MAX_TABLE_ROWS) -> list[list[object]]:
    """Extract rows from an openpyxl sheet while forward-filling merged cell values."""
    merged_map: dict[tuple[int, int], object] = {}
    merged_ranges = getattr(getattr(sheet, "merged_cells", None), "ranges", ())
    if merged_ranges:
        for rng in merged_ranges:
            try:
                top_left = sheet.cell(rng.min_row, rng.min_col).value
                if top_left is not None:
                    for r in range(rng.min_row, rng.max_row + 1):
                        for c in range(rng.min_col, rng.max_col + 1):
                            merged_map[(r, c)] = top_left
            except Exception:
                continue

    rows: list[list[object]] = []
    total_rows = getattr(sheet, "max_row", 0) or 0
    row_iter = sheet.iter_rows(values_only=False) if hasattr(sheet, "iter_rows") else ()

    truncated = False
    for r_idx, row in enumerate(row_iter, start=1):
        if len(rows) >= max_rows:
            # `max_row` counts blank rows too, so only an actual cut says so.
            truncated = True
            break
        row_vals: list[object] = []
        has_val = False
        for c_idx, cell in enumerate(row, start=1):
            val = merged_map.get((r_idx, c_idx), getattr(cell, "value", cell))
            if val is not None and str(val).strip():
                has_val = True
            row_vals.append(val)
        if has_val:
            rows.append(row_vals)

    if truncated:
        rows.append([f"... [Truncated: {total_rows} total rows, first {max_rows} rows shown] ..."])

    return rows


def _xlsx_rows(path: Path) -> list[tuple[str, list[list[object]]]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise RuntimeError("XLSX ingestion requires the 'office' dependency extra") from exc

    workbook = None
    if path.stat().st_size <= _FULL_LOAD_MAX_BYTES and _declared_merged_area(path) <= _MAX_MERGED_CELLS:
        try:
            workbook = load_workbook(path, data_only=True)
        except Exception:
            workbook = None
    if workbook is None:
        workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        results: list[tuple[str, list[list[object]]]] = []
        for sheet in workbook.worksheets:
            rows = _extract_sheet_rows_with_merged_cells(sheet)
            results.append((sheet.title, rows))
        return results
    finally:
        workbook.close()


def _xls_rows(path: Path) -> list[tuple[str, list[list[object]]]]:
    try:
        import pandas as pd  # type: ignore
    except ImportError as exc:
        raise RuntimeError("XLS ingestion requires the 'office' dependency extra") from exc
    try:
        frames = pd.read_excel(path, sheet_name=None, header=None)
    except ImportError as exc:
        raise RuntimeError("XLS ingestion requires xlrd from the 'office' dependency extra") from exc

    results: list[tuple[str, list[list[object]]]] = []
    for name, frame in frames.items():
        sheet_rows = frame.fillna("").values.tolist()
        total_rows = len(sheet_rows)
        if total_rows > MAX_TABLE_ROWS:
            sheet_rows = sheet_rows[:MAX_TABLE_ROWS]
            sheet_rows.append([f"... [Truncated: {total_rows} total rows, first {MAX_TABLE_ROWS} rows shown] ..."])
        results.append((str(name), sheet_rows))
    return results


def _docling_markdown(path: Path) -> str:
    try:
        # pyrefly: ignore [missing-import]
        from docling.document_converter import DocumentConverter  # type: ignore
    except ImportError:
        return ""
    try:
        return str(DocumentConverter().convert(str(path)).document.export_to_markdown() or "").strip()
    except Exception:
        return ""


def _archive_pages(path: Path, suffix: str) -> list[str]:
    prefix = _ARCHIVE_PREFIXES[suffix]["pages"]
    with zipfile.ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name.startswith(prefix) and name.endswith(".xml"))
        texts: list[str] = []
        for name in names:
            root = ElementTree.fromstring(archive.read(name))
            page_text = "\n".join(value.strip() for node in root.iter() if (value := node.text) and value.strip())
            if page_text:
                texts.append(page_text)
        return texts


def _archive_images(path: Path, document: EvidenceDocument) -> tuple[ImageBlock, ...]:
    prefix = _ARCHIVE_PREFIXES[path.suffix.lower()]["media"]
    with zipfile.ZipFile(path) as archive:
        page_by_name = _ppt_image_pages(archive) if path.suffix.lower() == _PPTX else {}
        names = sorted(name for name in archive.namelist() if name.startswith(prefix) and not name.endswith("/"))
        return tuple(
            ImageBlock(
                image_id=_id(document, "image", index),
                page=page_by_name.get(Path(name).name, 1),
                filename=Path(name).name,
                media_type=mimetypes.guess_type(name)[0] or "application/octet-stream",
                data=archive.read(name),
            )
            for index, name in enumerate(names, start=1)
        )


def _ppt_image_pages(archive: zipfile.ZipFile) -> dict[str, int]:
    pages: dict[str, int] = {}
    relationship_names = sorted(
        name for name in archive.namelist() if name.startswith("ppt/slides/_rels/slide") and name.endswith(".xml.rels")
    )
    for relationship_name in relationship_names:
        match = re.search(r"slide(\d+)\.xml\.rels$", relationship_name)
        if not match:
            continue
        page = int(match.group(1))
        root = ElementTree.fromstring(archive.read(relationship_name))
        for node in root.iter():
            target = str(node.attrib.get("Target", "") or "")
            if "/media/" in target or target.startswith("../media/"):
                pages.setdefault(Path(target).name, page)
    return pages


_TABLE_SEPARATOR_PATTERN = re.compile(r"^\|(?:\s*:?-[-:]*\s*\|)+$")


def _rows_to_markdown(rows: list[list[object]], sheet_name: str | None = None) -> str:
    truncation_note = None
    clean_rows = []
    for row in rows:
        if len(row) == 1 and str(row[0]).startswith("... [Truncated:"):
            truncation_note = str(row[0])
        else:
            clean_rows.append(row)

    normalized = [[_cell(value) for value in row] for row in clean_rows]
    normalized = [row for row in normalized if any(row)]
    if not normalized:
        return ""
    width = max(len(row) for row in normalized)
    padded = [row + [""] * (width - len(row)) for row in normalized]

    header_idx = 0
    title_lines: list[str] = []
    if len(padded) > 1 and width >= 2:
        # Check first up to 5 rows to locate best header row
        best_score = -1
        for idx in range(min(5, len(padded) - 1)):
            row = padded[idx]
            non_empty = [c for c in row if c.strip()]
            num_non_empty = len(non_empty)
            distinct_vals = len(set(non_empty))

            # If row has only 1 non-empty cell while width >= 3, it's a title banner
            if num_non_empty == 1 and width >= 3:
                continue

            # If row only contains 1 unique value across all cells (e.g. forward-filled merged banner), it's a title banner
            if distinct_vals <= 1 and width >= 2:
                continue

            # If row is mostly numeric, it is data, NOT column headers
            numeric_cells = sum(1 for c in non_empty if re.match(r"^[\$￥€£]?\s*-?\d+(?:[.,]\d+)?%?$", c.strip()))
            if num_non_empty > 0 and (numeric_cells / num_non_empty) >= 0.5:
                continue

            score = num_non_empty + distinct_vals
            if score > best_score:
                best_score = score
                header_idx = idx

        for idx in range(header_idx):
            non_empty = [c for c in padded[idx] if c.strip()]
            if non_empty:
                unique_vals = list(dict.fromkeys(non_empty))
                title_lines.append(" ".join(unique_vals))

    raw_header = padded[header_idx]
    header = [col if col.strip() else f"Column_{c_i}" for c_i, col in enumerate(raw_header, start=1)]
    body = padded[header_idx + 1 :]
    table_lines = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"] + [
        "| " + " | ".join(row) + " |" for row in body
    ]
    table_text = "\n".join(table_lines)
    if truncation_note:
        table_text += f"\n\n_{truncation_note}_"
    if title_lines:
        prefix_title = "\n\n".join(f"**{t}**" for t in title_lines)
        table_text = f"{prefix_title}\n\n{table_text}"
    if sheet_name:
        return f"### Sheet: {sheet_name}\n\n{table_text}"
    return table_text


def _flush_table_block(
    current_table: list[str],
    tables: list[TableBlock],
    start_index: int,
    document: EvidenceDocument,
    page: int,
    sheet: str | None,
) -> None:
    if len(current_table) >= 2 and any(_TABLE_SEPARATOR_PATTERN.match(row) for row in current_table):
        tbl_md = "\n".join(current_table)
        table_idx = start_index + len(tables)
        tables.append(
            TableBlock(
                table_id=_id(document, "table", table_idx),
                page=page,
                sheet=sheet,
                markdown=tbl_md,
            )
        )


def _extract_markdown_tables(
    markdown: str,
    document: EvidenceDocument,
    page: int = 1,
    start_index: int = 1,
) -> tuple[TableBlock, ...]:
    lines = markdown.splitlines()
    tables: list[TableBlock] = []
    current_table: list[str] = []
    current_sheet: str | None = None
    last_title: str | None = None

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("### Sheet:"):
            current_sheet = stripped[len("### Sheet:") :].strip()
            last_title = current_sheet
        elif stripped.startswith("### Table:"):
            current_sheet = stripped[len("### Table:") :].strip()
            last_title = current_sheet
        elif stripped.startswith("#"):
            last_title = stripped.lstrip("#").strip()

        if stripped.startswith("|") and stripped.endswith("|"):
            current_table.append(stripped)
        else:
            _flush_table_block(current_table, tables, start_index, document, page, current_sheet or last_title)
            current_table = []

    _flush_table_block(current_table, tables, start_index, document, page, current_sheet or last_title)
    return tuple(tables)


extract_markdown_tables = _extract_markdown_tables


def _cell(value: object) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip().replace("|", "\\|")


def _id(document: EvidenceDocument, kind: str, index: int) -> str:
    seed = f"{document.document_id}|{document.version}|{kind}|{index}"
    return f"{kind}-{hashlib.sha256(seed.encode('utf-8')).hexdigest()[:20]}"


__all__ = [
    "MAX_TABLE_ROWS",
    "OFFICE_EXTENSIONS",
    "extract_markdown_tables",
    "load_office_document",
    "parsed_to_documents",
]
